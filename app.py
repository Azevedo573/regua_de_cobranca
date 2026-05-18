"""
Sistema de Processamento de Faturamento
Gera planilhas filtradas por vencimento e status para disparo WhatsApp
"""
import streamlit as st
import pandas as pd
from datetime import datetime, timedelta
import io
from openpyxl.styles import numbers
import os
from dotenv import load_dotenv
import time

# Carrega variáveis de ambiente
load_dotenv()

# Importa integração WhatsApp
try:
    from whatsapp_simples import WhatsAppClientSimples
except ImportError:
    st.warning("⚠️ whatsapp_simples.py não encontrado. Funcionalidade WhatsApp desabilitada.")

# Configuração da página
st.set_page_config(
    page_title="Gerador de Planilhas de Cobrança",
    page_icon="📊",
    layout="wide"
)

# Título principal
st.title("📊 Gerador de Planilhas de Cobrança")
st.markdown("---")

def calcular_dias_vencimento(row):
    """Calcula dias desde o vencimento (positivo = vencido, negativo = a vencer)"""
    data_venc = row.get('Data de Vencimento')
    if pd.isna(data_venc):
        return None
    
    # Converte para datetime se necessário
    if isinstance(data_venc, str):
        try:
            data_venc = pd.to_datetime(data_venc, dayfirst=True)
        except:
            return None
    
    hoje = datetime.now().date()
    if hasattr(data_venc, 'date'):
        data_venc = data_venc.date()
    
    return (hoje - data_venc).days



def gerar_excel(df, nome_arquivo):
    """Gera arquivo Excel em memória para download"""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Dados')
        # Formata coluna de data
        if 'Data de Vencimento' in df.columns:
            worksheet = writer.sheets['Dados']
            col_num = df.columns.get_loc('Data de Vencimento') + 1
            for row in range(2, len(df) + 2):
                cell = worksheet.cell(row=row, column=col_num)
                cell.number_format = 'dd/mm/yyyy'
    output.seek(0)
    return output

# Função para gerar planilha resumida
def gerar_excel_resumido(df):
    colunas_resumidas = [
        'Celular',
        'Nome Titular',
        'Operadora da Fatura',
        'Data de Vencimento',
        'Valor Cobrança',
        'Alerta'
    ]
    # Garante que só as colunas existentes serão usadas
    colunas_existentes = [c for c in colunas_resumidas if c in df.columns]
    df_resumido = df[colunas_existentes].copy() if colunas_existentes else pd.DataFrame()
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_resumido.to_excel(writer, index=False, sheet_name='Resumido')
        # Formata coluna de data
        if 'Data de Vencimento' in df_resumido.columns:
            worksheet = writer.sheets['Resumido']
            col_num = df_resumido.columns.get_loc('Data de Vencimento') + 1
            for row in range(2, len(df_resumido) + 2):
                cell = worksheet.cell(row=row, column=col_num)
                cell.number_format = 'dd/mm/yyyy'
    output.seek(0)
    return output

# ============================================================================
# FUNÇÕES PARA INTEGRAÇÃO WHATSAPP SIMPLES
# ============================================================================

def inicializar_whatsapp_simples():
    """Inicializa cliente WhatsApp simples"""
    try:
        api_url = os.getenv('WHATSAPP_API_URL')
        auth_token = os.getenv('WHATSAPP_AUTH_TOKEN')
        template_id = os.getenv('WHATSAPP_TEMPLATE_COBRANCA')
        # template_id_1 = os.getenv('WHATSAPP_TEMPLATE_1')

        
        if not api_url or not auth_token or not template_id:
            return None, None
        
        client = WhatsAppClientSimples(
            api_url=api_url,
            auth_token=auth_token
        )
        
        return client, template_id
    except Exception as e:
        st.error(f"❌ Erro ao inicializar WhatsApp: {str(e)}")
        return None, None

def enviar_whatsapp_simples(df, client, template_id, periodo, simular=False):
    """
    Envia mensagens WhatsApp para todos os números da lista com template automático por período.
    
    Espera coluna 'Celular' com os telefones
    
    Args:
        df: DataFrame com os dados
        client: Cliente WhatsApp
        template_id: ID da template (ignorado, usa automático baseado no período)
        periodo: Nome do período (ex: "1-3 dias de atraso")
        simular: Se True, não envia de verdade
    """
    
    if df is None or df.empty:
        st.error("❌ Nenhum registro para enviar")
        return None
    
    # Verifica coluna de telefone
    if 'Celular' not in df.columns:
        st.error("❌ Planilha não possui coluna 'Celular'")
        return None
    
    # Prepara lista de contatos
    contatos = []
    
    for idx, row in df.iterrows():
        telefone = str(row.get('Celular', '')).strip()
        
        # Pula se não tem telefone
        if not telefone or telefone == 'nan' or telefone == '':
            continue
        
        nome = str(row.get('Nome Titular', 'Cliente')).strip()
        
        contatos.append({
            'nome': nome,
            'telefone': telefone,
            'valor': str(row.get('Valor Cobrança', '0')).strip(),
            'vencimento': str(row.get('Data de Vencimento', '')).strip(),
            'operadora': str(row.get('Operadora da Fatura', '')).strip(),
        })
    
    if not contatos:
        st.error("❌ Nenhum contato com telefone encontrado")
        return None
    
    st.info(f"📤 Preparados {len(contatos)} contatos para envio")
    
    # Define função para extrair parâmetros baseado no período
    def extrair_parametros(contato):
        """
        Extrai parâmetros diferentes baseado no tipo de template do período.
        Cada template pode ter um número diferente de parâmetros.
        """
        tipo_template = client.obter_tipo_periodo(periodo)
        
        # Template ANTECIPADO: Nome, Operadora
        if tipo_template == "antecipado":
            return [
                contato['nome'],          # {{0}} - Nome
                contato['operadora']      # {{1}} - Operadora
            ]
        
        # Template A_VENCER: Nome, Operadora
        elif tipo_template == "a_vencer":
            return [
                contato['nome'],          # {{0}} - Nome
                contato['operadora']      # {{1}} - Operadora
            ]
        
        # Template ATRASO_LEVE: Nome, Operadora, Data
        elif tipo_template == "atraso_leve":
            return [
                contato['nome'],          # {{0}} - Nome
                contato['operadora'],     # {{1}} - Operadora
                contato['vencimento']     # {{2}} - Data de Vencimento
            ]
        
        # Template ATRASO_MODERADO: Nome, Operadora, Valor
        elif tipo_template == "atraso_moderado":
            return [
                contato['nome'],          # {{0}} - Nome
                contato['operadora'],     # {{1}} - Operadora
                contato['valor']          # {{2}} - Valor
            ]
        
        # Template ATRASO_SEVERO: Nome, Operadora, Valor
        elif tipo_template == "atraso_severo":
            return [
                contato['nome'],          # {{0}} - Nome
                contato['operadora'],     # {{1}} - Operadora
                contato['valor']          # {{2}} - Valor
            ]
        
        # DEFAULT: Nome, Operadora (mais seguro)
        else:
            return [
                contato['nome'],
                contato['operadora']
            ]
    
    # Envia lote com template AUTOMÁTICO baseado no período
    with st.spinner("📤 Enviando mensagens..."):
        resultado = client.enviar_lote_por_periodo(
            contatos=contatos,
            periodo=periodo,                          # Template automático!
            funcao_parametros=extrair_parametros,
            simular=simular,
            intervalo=1.5
        )
    
    return resultado

def exibir_resultado_whatsapp_simples(resultado):
    """Exibe resultado do envio WhatsApp com detalhes de WAMID e período"""
    if resultado is None:
        return
    
    # Exibe período e template (se disponível)
    if 'periodo' in resultado and 'tipo_template' in resultado:
        col1, col2 = st.columns(2)
        with col1:
            st.info(f"📄 **Período:** {resultado['periodo']}")
        with col2:
            st.info(f"📱 **Template:** {resultado['tipo_template'].upper().replace('_', ' ')}")
        st.markdown("")
    
    # Exibe estatísticas
    col1, col2, col3, col4 = st.columns(4)
    
    with col1:
        st.metric("📊 Total", resultado['total'])
    with col2:
        st.metric("✅ Sucesso", resultado['sucesso'])
    with col3:
        st.metric("❌ Erro", resultado['erro'])
    with col4:
        st.metric("📈 Taxa", resultado['taxa_sucesso'])
    
    st.markdown("---")
    
    # Mostra detalhes de cada envio
    st.subheader("📋 Detalhes do Envio")
    
    # Sucesso
    sucessos = [r for r in resultado['resultados'] if r['sucesso']]
    if sucessos:
        with st.expander(f"✅ Enviadas ({len(sucessos)})", expanded=True):
            for r in sucessos:
                col_nome, col_telefone, col_wamid = st.columns([2, 2, 3])
                with col_nome:
                    st.write(r['nome'])
                with col_telefone:
                    st.write(f"`{r['telefone']}`")
                with col_wamid:
                    if r['wamid'] and r['wamid'] != 'simulado':
                        st.write(f"🔗 `{r['wamid'][:30]}...`")
                    else:
                        st.write("(Simulado)")
    
    # Falhas
    falhas = [r for r in resultado['resultados'] if not r['sucesso']]
    if falhas:
        with st.expander(f"❌ Falhadas ({len(falhas)})"):
            for r in falhas:
                col_nome, col_telefone, col_erro = st.columns([2, 2, 3])
                with col_nome:
                    st.write(r['nome'])
                with col_telefone:
                    st.write(f"`{r['telefone']}`")
                with col_erro:
                    st.write(f"{r['mensagem']}")

def main():
    # Inicializa session_state para armazenar resultados
    if "resultados" not in st.session_state:
        st.session_state.resultados = None
    if "df_original" not in st.session_state:
        st.session_state.df_original = None
    if "ultimo_arquivo" not in st.session_state:
        st.session_state.ultimo_arquivo = None
    
    # Sidebar
    st.sidebar.header("📋 Instruções")
    st.sidebar.markdown("""
    1. Faça upload da planilha de faturamento
    2. Sistema filtra automaticamente
    3. Baixe as planilhas filtradas
    4. Clique em "📤 Disparar WhatsApp" para enviar
    """)
    
    st.sidebar.markdown("---")
    st.sidebar.header("📁 Colunas Esperadas")
    st.sidebar.markdown("""
    **Obrigatórias:**
    - Data de Vencimento
    - Celular
    - Nome Titular
    
    **Opcionais:**
    - Valor Cobrança
    - Status da Fatura
    """)
    
    # Inicializa WhatsApp
    st.sidebar.markdown("---")
    st.sidebar.header("📱 WhatsApp")
    
    client, template_id = inicializar_whatsapp_simples()
    
    if client:
        ok, msg = client.testar_conexao()
        if ok:
            st.sidebar.success("✅ API conectada")
        else:
            st.sidebar.error(f"❌ {msg}")
    else:
        st.sidebar.warning("⚠️ Configure .env")
    
    # Modo teste
    modo_simulacao = st.sidebar.checkbox(
        "🧪 Modo Simulação",
        value=False,
        help="Simula sem enviar de verdade"
    )
    
    # Upload
    uploaded_file = st.file_uploader(
        "📤 Arraste ou selecione a planilha",
        type=['xlsx', 'xls']
    )
    
    if uploaded_file is not None:
        try:
            # Verifica se arquivo mudou ou se precisa processar
            arquivo_mudou = (st.session_state.ultimo_arquivo != uploaded_file.name)
            
            if arquivo_mudou or st.session_state.resultados is None:
                # Cria placeholders para feedback
                status_placeholder = st.empty()
                progress_bar = st.progress(0)
                
                # Carrega arquivo
                status_placeholder.info("📥 Carregando arquivo...")
                df = pd.read_excel(uploaded_file)
                st.session_state.df_original = df
                st.session_state.ultimo_arquivo = uploaded_file.name
                progress_bar.progress(10)
                
                # Verifica coluna obrigatória
                if 'Data de Vencimento' not in df.columns:
                    status_placeholder.error("❌ Coluna 'Data de Vencimento' não encontrada!")
                    progress_bar.empty()
                    status_placeholder.empty()
                    return
                
                st.success(f"✅ {len(df)} registros carregados")
                progress_bar.progress(15)
                
                # Preview
                with st.expander("👀 Visualizar dados"):
                    st.dataframe(df.head(10))
                
                st.markdown("---")
                st.header("📑 Planilhas por Período")
                
                # Filtra dados
                status_placeholder.info("⚙️ Filtrando dados...")
                progress_bar.progress(20)
                
                filtros = [
                    {"nome": "Cobrança Antecipada - 10 dias", "dias_min": -10, "dias_max": -10},
                    {"nome": "Cobrança Antecipada - 5 dias", "dias_min": -5, "dias_max": -5},
                    {"nome": "Cobrança Antecipada - 3 dias", "dias_min": -3, "dias_max": -3},
                    {"nome": "A vencer (Até hoje)", "dias_min": None, "dias_max": 0},
                    {"nome": "A vencer (hoje)", "dias_min": 0, "dias_max": 0},
                    {"nome": "1-3 dias de atraso", "dias_min": 1, "dias_max": 3},
                    {"nome": "4-7 dias de atraso", "dias_min": 4, "dias_max": 7},
                    {"nome": "8-14 dias de atraso", "dias_min": 8, "dias_max": 14},
                    {"nome": "15-30 dias de atraso", "dias_min": 15, "dias_max": 30},
                    {"nome": "31-60 dias de atraso", "dias_min": 31, "dias_max": 60},
                    {"nome": "61-90 dias de atraso", "dias_min": 61, "dias_max": 90},
                    {"nome": "> 90 dias de atraso", "dias_min": 91, "dias_max": None},
                ]
                
                # Processa cada filtro
                resultados = {}
                
                for idx, filtro in enumerate(filtros):
                    df_filtro = df.copy()
                    df_filtro['Dias_Vencimento'] = df_filtro.apply(calcular_dias_vencimento, axis=1)
                    df_filtro = df_filtro.dropna(subset=['Dias_Vencimento'])
                    
                    # Aplica filtro
                    dias_min = filtro['dias_min']
                    dias_max = filtro['dias_max']
                    
                    if dias_min is None and dias_max is not None:
                        df_resultado = df_filtro[df_filtro['Dias_Vencimento'] <= dias_max]
                    elif dias_min is not None and dias_max is None:
                        df_resultado = df_filtro[df_filtro['Dias_Vencimento'] >= dias_min]
                    elif dias_min is not None and dias_max is not None:
                        df_resultado = df_filtro[
                            (df_filtro['Dias_Vencimento'] >= dias_min) & 
                            (df_filtro['Dias_Vencimento'] <= dias_max)
                        ]
                    else:
                        df_resultado = pd.DataFrame()
                    
                    # Filtra status pagos
                    if 'Status da Fatura' in df.columns:
                        status_pagos = ['Pago', 'PAGO', 'pago', 'Paga', 'PAGA', 'paga']
                        status_canceladas = ['Cancelada', 'CANCELADA', 'cancelada']
                        status_baixadas = ['Baixada', 'BAIXADA', 'baixada']
                        
                        df_resultado = df_resultado[
                            (~df_resultado['Status da Fatura'].isin(status_pagos)) &
                            (~df_resultado['Status da Fatura'].isin(status_canceladas)) &
                            (~df_resultado['Status da Fatura'].isin(status_baixadas)) &
                            (df_resultado['Status da Fatura'].notna())
                        ]
                    
                    # Remove coluna auxiliar
                    if 'Dias_Vencimento' in df_resultado.columns:
                        df_resultado = df_resultado.drop(columns=['Dias_Vencimento'])
                    
                    resultados[filtro['nome']] = df_resultado
                    
                    # Atualiza progress bar
                    progresso = int(20 + (idx / len(filtros)) * 70)
                    progress_bar.progress(progresso)
                
                # Salva resultados em session_state
                st.session_state.resultados = resultados
                
                status_placeholder.success("✅ Processamento concluído!")
                progress_bar.progress(100)
                
                # Limpa placeholders após 1 segundo
                time.sleep(1)
                status_placeholder.empty()
                progress_bar.empty()
            
            # =========================================================================
            # EXIBE RESULTADOS (reutiliza dados do session_state)
            # =========================================================================
            
            if st.session_state.resultados is not None:
                st.markdown("---")
                st.header("📊 Resultados Processados")
                
                # Exibe informações sobre templates
                if client:
                    with st.expander("📱 ℹ️ Templates de WhatsApp Disponíveis"):
                        col1, col2 = st.columns(2)
                        
                        templates = client.listar_templates()
                        for i, (tipo, config) in enumerate(templates.items()):
                            if i % 2 == 0:
                                col = col1
                            else:
                                col = col2
                            
                            with col:
                                st.markdown(f"**{tipo.upper().replace('_', ' ')}**")
                                st.caption(f"Template ID: `{config['id']}`")
                                st.caption(f"🔘 Botões: {', '.join(config['botoes'])}")
                
                st.markdown("")  # Espaço
                
                resultados = st.session_state.resultados
                total_geral = 0
                
                for nome, df_resultado in resultados.items():
                    with st.expander(f"📄 {nome} ({len(df_resultado)} registros)"):
                        if len(df_resultado) > 0:
                            # Métricas
                            col1, col2, col3 = st.columns(3)
                            with col1:
                                st.metric("Registros", len(df_resultado))
                            with col2:
                                if 'Valor Cobrança' in df_resultado.columns:
                                    valor = df_resultado['Valor Cobrança'].sum()
                                    st.metric("Valor Total", f"R$ {valor:,.2f}")
                                else:
                                    st.metric("Valor Total", "N/A")
                            with col3:
                                com_cel = (df_resultado['Celular'].notna()).sum() if 'Celular' in df_resultado.columns else 0
                                st.metric("Com Celular", com_cel)
                            
                            # Preview
                            st.dataframe(df_resultado.head(5), width='stretch')
                            
                            # Downloads
                            excel_data = gerar_excel(df_resultado, f"{nome}.xlsx")
                            st.download_button(
                                label=f"⬇️ Baixar {nome}",
                                data=excel_data,
                                file_name=f"{nome.replace(' ', '_')}.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                key=f"download_{nome}"
                            )
                            
                            excel_resumido = gerar_excel_resumido(df_resultado)
                            st.download_button(
                                label=f"⬇️ Baixar {nome} (Resumido)",
                                data=excel_resumido,
                                file_name=f"{nome.replace(' ', '_')}_resumido.xlsx",
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                key=f"download_resumido_{nome}"
                            )
                            
                            # Botão WhatsApp
                            if client and template_id and len(df_resultado) > 0:
                                st.markdown("---")
                                
                                # Exibe tipo de template que será usado
                                tipo_template = client.obter_tipo_periodo(nome)
                                config_template = client.obter_template_por_periodo(nome)
                                botoes = list(config_template['response_action']['buttonActions'].keys())
                                
                                col_info1, col_info2 = st.columns(2)
                                with col_info1:
                                    st.info(f"📱 **Tipo de Template:** {tipo_template.upper().replace('_', ' ')}")
                                with col_info2:
                                    st.info(f"🔘 **Botões:** {', '.join(botoes)}")
                                
                                st.markdown("")
                                
                                if st.button(
                                    f"📤 Disparar {len(df_resultado)} no WhatsApp",
                                    key=f"btn_whatsapp_{nome}",
                                    type="primary"
                                ):
                                    # Preview antes de enviar
                                    with st.expander("👀 Ver contatos que serão enviados"):
                                        df_preview = df_resultado[['Celular', 'Nome Titular']].head(10)
                                        st.dataframe(df_preview)
                                    
                                    # Envia com template automático do período
                                    resultado = enviar_whatsapp_simples(
                                        df_resultado,
                                        client,
                                        template_id,
                                        periodo=nome,                  # Nome do período!
                                        simular=modo_simulacao
                                    )
                                    
                                    if resultado:
                                        exibir_resultado_whatsapp_simples(resultado)
                            
                            total_geral += len(df_resultado)
                        else:
                            st.info("Nenhum registro encontrado")
                
                st.markdown("---")
                st.write(f"✅ Total filtrado: **{total_geral}**")
        
        except Exception as e:
            st.error(f"❌ Erro: {str(e)}")
            st.info("Verifique o formato da planilha")
    
    else:
        st.info("👆 Faça upload de uma planilha Excel para começar")
        col1, col2, col3 = st.columns(3)
        with col3:
            st.markdown("""
            **Opcionais:**
            - Status da Fatura
            - Valor Cobrança
            """)

if __name__ == "__main__":
    main()